# -*- coding: utf-8 -*-
"""LocalFileReaderService 전체 확장자 테스트

지원 확장자: TXT, MD, PDF, DOCX, XLSX, PPTX, HWP
각 파일 1000자 이상 테스트 데이터 생성 후 search / read_chunk 검증
"""
import os
import sys
import json
import tempfile
import shutil

# 프로젝트 루트 → python/ 경로 추가
PYTHON_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if PYTHON_DIR not in sys.path:
    sys.path.insert(0, PYTHON_DIR)

from services.local_file_reader_service import LocalFileReaderService

# ============================================================================
# 테스트용 텍스트 (1200자 이상)
# ============================================================================
SAMPLE_TEXT = """인서티(Inserty)는 한글(HWP) 문서 편집을 돕는 AI 데스크톱 애플리케이션입니다.
소상공인과 중소기업이 지원사업 서류를 작성할 때 가장 큰 병목은 HWP 양식 문서의 표·셀 구조를 이해하고
각 항목에 맞는 내용을 정확히 채워 넣는 작업입니다. Inserty는 이 과정을 자동화하여 작성 시간을 크게 단축합니다.

주요 기능:
1. 양식쌍 기반 작성: 빈 양식(템플릿)과 작성된 예시를 학습하여 새 문서 작성을 지원합니다.
2. 참조 자료 활용: 프로젝트 파일과 채팅 업로드 파일을 검색하여 맥락에 맞는 내용을 제안합니다.
3. 실시간 편집: 사용자 요청에 따라 문서를 직접 수정하며, Track Changes 기능으로 모든 변경 사항을 추적합니다.
4. Multi-Document 지원: 여러 HWP 문서를 동시에 관리하고 선택적으로 편집할 수 있습니다.

기술 스택:
- Frontend: React 18.3 + TypeScript 5.4 + Vite 5.4 + Tailwind CSS 3.4
- Desktop: Electron 33.2 (Main + Preload + Renderer)
- Backend: Python 3.12 + pyhwpx (HWP COM 자동화)
- State Management: Zustand 4.5
- Database: better-sqlite3 (로컬 SQLite)
- IPC: JSON-RPC over stdio (Electron ↔ Python)

사업화 전략:
Inserty는 B2C 개인 구독 모델에서 시작하여, 향후 B2B 채널(지역 센터, 협회, 컨설턴트 조직)로 확장할 계획입니다.
초기 목표는 반송시장과 협력 컨설턴트를 통한 데모·체험·유료 전환 실험이며,
2026년 내 초기 고객 2,000명 확보 및 유료 전환 100건을 달성하는 것을 목표로 합니다.

팀 구성:
- 대표 이학빈: 제품 방향, 문서 파이프라인 설계, 프로덕트 총괄
- 팀원 양온유: S/W 개발, 데이터 엔지니어, LLM 디버깅

핵심 차별점:
기존 범용 AI 도구와 달리, Inserty는 HWP 문서에 직접 연결하여 셀 단위로 작성하고,
RAG 기반으로 할루시네이션을 최소화하며, 전체/부분 승인·거절 UX를 제공합니다.
작성 시간 70% 단축, 오류 감소, 전환율·유지율 지표를 KPI로 관리합니다.
"""

# 검색 키워드 (모든 확장자에 공통)
SEARCH_QUERIES = ["팀 구성", "기술 스택", "사업화 전략"]
UNIQUE_KEYWORD = "반송시장"  # 정확히 1곳에만 등장


def _print_header(title: str) -> None:
    print(f"\n{'=' * 60}")
    print(f"  {title}")
    print(f"{'=' * 60}")


def _print_result(label: str, ok: bool, detail: str = "") -> None:
    status = "✅ PASS" if ok else "❌ FAIL"
    print(f"  [{status}] {label}" + (f" — {detail}" if detail else ""))


# ============================================================================
# 테스트 데이터 생성
# ============================================================================

def create_txt(dirpath: str) -> str:
    path = os.path.join(dirpath, "test_sample.txt")
    with open(path, "w", encoding="utf-8") as f:
        f.write(SAMPLE_TEXT)
    return path


def create_md(dirpath: str) -> str:
    path = os.path.join(dirpath, "test_sample.md")
    with open(path, "w", encoding="utf-8") as f:
        f.write("# Inserty AI 소개\n\n" + SAMPLE_TEXT)
    return path


def create_pdf(dirpath: str) -> str:
    path = os.path.join(dirpath, "test_sample.pdf")
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.pdfgen import canvas
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont

        c = canvas.Canvas(path, pagesize=A4)
        # 한글 폰트 등록
        font_name = "Helvetica"
        for fp in [
            "C:/Windows/Fonts/malgun.ttf",
            "C:/Windows/Fonts/gulim.ttc",
        ]:
            if os.path.exists(fp):
                try:
                    pdfmetrics.registerFont(TTFont("KoreanFont", fp))
                    font_name = "KoreanFont"
                    break
                except Exception:
                    pass

        c.setFont(font_name, 10)
        y = 780
        for line in SAMPLE_TEXT.split("\n"):
            if y < 50:
                c.showPage()
                c.setFont(font_name, 10)
                y = 780
            c.drawString(40, y, line[:90])  # 줄 폭 제한
            y -= 14
        c.save()
        return path
    except ImportError:
        print("  ⚠️  reportlab 없음 → PDF 테스트 스킵")
        return ""


def create_docx(dirpath: str) -> str:
    path = os.path.join(dirpath, "test_sample.docx")
    try:
        from docx import Document
        doc = Document()
        doc.add_heading("Inserty AI 소개", level=1)
        for para in SAMPLE_TEXT.split("\n"):
            if para.strip():
                doc.add_paragraph(para.strip())
        doc.save(path)
        return path
    except ImportError:
        print("  ⚠️  python-docx 없음 → DOCX 테스트 스킵")
        return ""


def create_xlsx(dirpath: str) -> str:
    path = os.path.join(dirpath, "test_sample.xlsx")
    try:
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "Inserty"
        for i, line in enumerate(SAMPLE_TEXT.split("\n"), 1):
            if line.strip():
                ws.cell(row=i, column=1, value=line.strip())
        wb.save(path)
        return path
    except ImportError:
        print("  ⚠️  openpyxl 없음 → XLSX 테스트 스킵")
        return ""


def create_pptx(dirpath: str) -> str:
    path = os.path.join(dirpath, "test_sample.pptx")
    try:
        from pptx import Presentation
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[1])  # title + content
        slide.shapes.title.text = "Inserty AI 소개"
        body = slide.placeholders[1]
        body.text = SAMPLE_TEXT[:2000]
        prs.save(path)
        return path
    except ImportError:
        print("  ⚠️  python-pptx 없음 → PPTX 테스트 스킵")
        return ""


# ============================================================================
# 테스트 실행
# ============================================================================

def test_read_file_text(reader: LocalFileReaderService, ext: str, path: str) -> bool:
    """read_file_text 단위 테스트: 텍스트 추출 + 최소 길이 확인"""
    text = reader.read_file_text(path)
    has_content = len(text) >= 100
    has_no_error = not text.startswith("[")  # 에러 메시지 체크
    ok = has_content and has_no_error
    detail = f"{len(text):,}자"
    if not has_content:
        detail += " (내용 부족!)"
    if not has_no_error:
        detail += f" ERROR: {text[:200]}"
    _print_result(f"read_file_text ({ext})", ok, detail)
    if ok:
        # 키워드 포함 확인
        has_keyword = UNIQUE_KEYWORD in text
        _print_result(f"  키워드 '{UNIQUE_KEYWORD}' 포함", has_keyword,
                      f"{'found' if has_keyword else 'NOT FOUND'}")
    return ok


def test_search(reader: LocalFileReaderService, ext: str, path: str,
                file_info: dict) -> bool:
    """search() 단위 테스트: 키워드 검색 + start_char/end_char 검증"""
    # file_name 미지정 (키워드 검색)
    result = reader.search(SEARCH_QUERIES, project_id="test_proj")
    ok = result.get("success") and len(result.get("results", [])) > 0
    count = len(result.get("results", []))
    _print_result(f"search keywords ({ext})", ok, f"{count}개 매치")

    if ok:
        # start_char / end_char 검증
        r0 = result["results"][0]
        has_positions = "start_char" in r0 and "end_char" in r0
        _print_result(f"  start_char/end_char 존재", has_positions,
                      f"start={r0.get('start_char')}, end={r0.get('end_char')}")

    # file_name 지정 (해당 파일 내부 snippet 검색)
    result2 = reader.search([UNIQUE_KEYWORD], project_id="test_proj",
                            file_name=file_info["display_name"])
    ok2 = result2.get("success") and len(result2.get("results", [])) > 0
    if ok2:
        content = result2["results"][0].get("content", "")
        is_snippet = len(content) < len(reader.read_file_text(path))
        has_positions = "start_char" in result2["results"][0] and "end_char" in result2["results"][0]
        _print_result(
            f"search file-scoped snippet ({ext})",
            bool(is_snippet and has_positions),
            f"{len(content):,}자, start={result2['results'][0].get('start_char')}, end={result2['results'][0].get('end_char')}",
        )
    else:
        _print_result(f"search file-scoped snippet ({ext})", False,
                      str(result2.get("error", result2.get("message", ""))))
    return ok


def test_read_chunk(reader: LocalFileReaderService, ext: str,
                    file_info: dict) -> bool:
    """read_chunk() 단위 테스트: offset/length 구간 읽기"""
    result = reader.read_chunk(file_info["display_name"],
                               project_id="test_proj", offset=100, length=500)
    ok = result.get("success") and len(result.get("content", "")) > 0
    detail = ""
    if ok:
        detail = (f"offset={result['offset']}, len={result['length']}, "
                  f"total={result.get('total_chars', '?')}")
    else:
        detail = str(result.get("error", ""))
    _print_result(f"read_chunk ({ext})", ok, detail)
    return ok


def run_tests():
    """모든 확장자에 대해 테스트 실행"""
    # 임시 디렉토리 구조: tmp/projects/test_proj/
    tmpdir = tempfile.mkdtemp(prefix="inserty_rag_test_")
    proj_dir = os.path.join(tmpdir, "projects", "test_proj")
    os.makedirs(proj_dir, exist_ok=True)

    print(f"테스트 디렉토리: {tmpdir}")

    # 테스트 파일 생성
    _print_header("1. 테스트 데이터 생성")
    creators = {
        ".txt": create_txt,
        ".md": create_md,
        ".pdf": create_pdf,
        ".docx": create_docx,
        ".xlsx": create_xlsx,
        ".pptx": create_pptx,
    }

    created_files: dict = {}
    for ext, creator in creators.items():
        path = creator(proj_dir)
        if path:
            created_files[ext] = path
            size = os.path.getsize(path)
            print(f"  {ext}: {os.path.basename(path)} ({size:,} bytes)")

    # HWP 파일 복사
    hwp_src = r"C:\Users\dlgkr\Downloads\[신청서] 2026년 부니콘 씨드 육성사업(부산 예비창업패키지) (1) (2).hwp"
    if os.path.exists(hwp_src):
        hwp_dst = os.path.join(proj_dir, os.path.basename(hwp_src))
        shutil.copy2(hwp_src, hwp_dst)
        created_files[".hwp"] = hwp_dst
        print(f"  .hwp: {os.path.basename(hwp_dst)} ({os.path.getsize(hwp_dst):,} bytes)")
    else:
        print(f"  ⚠️  HWP 파일 없음: {hwp_src}")

    # DB 생성 (files 테이블)
    import sqlite3
    db_path = os.path.join(tmpdir, "inserty.db")
    conn = sqlite3.connect(db_path)
    conn.execute("""
        CREATE TABLE IF NOT EXISTS files (
            id TEXT PRIMARY KEY,
            scope TEXT NOT NULL DEFAULT 'project',
            project_id TEXT,
            chat_id TEXT,
            name TEXT NOT NULL,
            rel_path TEXT NOT NULL,
            extension TEXT,
            added_at TEXT DEFAULT (datetime('now'))
        )
    """)
    for i, (ext, path) in enumerate(created_files.items()):
        fname = os.path.basename(path)
        rel_path = fname  # proj_dir 기준 상대경로
        conn.execute(
            "INSERT INTO files (id, scope, project_id, name, rel_path, extension) VALUES (?,?,?,?,?,?)",
            (f"file_{i}", "project", "test_proj", fname, rel_path, ext),
        )
    conn.commit()
    conn.close()
    print(f"  DB: {db_path} ({len(created_files)}개 파일 등록)")

    # LocalFileReaderService 초기화
    reader = LocalFileReaderService(tmpdir)

    # 파일 목록 확인
    _print_header("2. get_available_files")
    files = reader.get_available_files("test_proj")
    print(f"  파일 수: {len(files)}")
    file_map: dict = {}
    for f in files:
        ext = f["extension"]
        file_map[ext] = f
        print(f"  {ext}: {f['display_name']} → {f['abs_path']}")
        exists = os.path.exists(f["abs_path"])
        _print_result(f"  파일 존재 확인", exists)

    # 확장자별 테스트
    results_summary: dict = {}

    for ext in [".txt", ".md", ".pdf", ".docx", ".xlsx", ".pptx", ".hwp"]:
        _print_header(f"3-{ext}: {ext.upper()} 파일 테스트")
        if ext not in file_map:
            print(f"  ⚠️  스킵 (파일 없음)")
            results_summary[ext] = "SKIP"
            continue

        fi = file_map[ext]

        # read_file_text
        read_ok = test_read_file_text(reader, ext, fi["abs_path"])

        # search (키워드 + 전문)
        search_ok = test_search(reader, ext, fi["abs_path"], fi)

        # read_chunk
        chunk_ok = test_read_chunk(reader, ext, fi)

        results_summary[ext] = "PASS" if (read_ok and search_ok and chunk_ok) else "FAIL"

    # 최종 요약
    _print_header("최종 결과 요약")
    all_pass = True
    for ext, status in results_summary.items():
        icon = "✅" if status == "PASS" else ("⚠️" if status == "SKIP" else "❌")
        print(f"  {icon} {ext}: {status}")
        if status == "FAIL":
            all_pass = False

    print(f"\n{'=' * 60}")
    if all_pass:
        print("  🎉 모든 테스트 통과!")
    else:
        print("  ⚠️  일부 테스트 실패 — 위 로그 확인")
    print(f"{'=' * 60}")

    # 정리
    try:
        shutil.rmtree(tmpdir)
        print(f"\n  🧹 임시 디렉토리 삭제 완료: {tmpdir}")
    except Exception as e:
        print(f"\n  ⚠️  임시 디렉토리 삭제 실패: {e}")


if __name__ == "__main__":
    run_tests()
