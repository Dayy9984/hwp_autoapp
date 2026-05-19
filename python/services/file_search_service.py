# -*- coding: utf-8 -*-



"""



File Search Service







Replaces LightRAG-based RAG with OpenAI Vector Stores + Responses API file_search.



This module keeps the existing public interface used by Electron/Python bridge.



"""







import json



import os
import sqlite3



import time



from datetime import datetime



from typing import Dict, Any, Optional, Callable, List, Tuple







from openai import OpenAI







from .config import config



from utils.logger import debug, log_file_search_storage, log_file_search_tool_usage











# OpenAI가 native로 지원하고 프로젝트에서 허용하는 확장자만



FILE_SEARCH_NATIVE_EXTS = {



    ".pdf",      # PDF 문서



    ".txt",      # 텍스트 파일



    ".md",       # Markdown



    ".docx",     # MS Word



    ".pptx",     # MS PowerPoint



}



# 변환이 필요한 확장자 (프로젝트 허용 파일 중)



CONVERT_TO_DOCX = {".hwp", ".hwpx"}  # 한글 확장자 식별(안전모드: text_content 업로드)



CONVERT_TO_MD = {".xls", ".xlsx", ".xlsm"}  # Excel → Markdown 표 변환



UNSUPPORTED_IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".gif", ".webp"}



TEXT_FALLBACK_EXT = ".txt"



POLL_INTERVAL_SEC = 1.0



POLL_TIMEOUT_SEC = 240.0

# 채팅 스코프 벡터 스토어 TTL (일)
CHAT_VECTOR_STORE_TTL_DAYS = int(os.getenv("CHAT_VECTOR_STORE_TTL_DAYS", "7"))
DB_FILENAME = "inserty.db"
LEGACY_DB_FILENAME = "mallo.db"
VECTOR_STORE_PREFIX = "inserty"











class FileSearchService:

    """OpenAI File Search ?? ??? (? RAG ??)."""

    def __init__(self, api_key: Optional[str] = None):
        self.openai_api_key = (api_key or "").strip()
        if not self.openai_api_key:
            raise ValueError("OPENAI_API_KEY_REQUIRED")

        self.max_network_retries = max(1, int(os.getenv("FILE_SEARCH_NETWORK_RETRIES", "3")))
        self.retry_backoff_sec = max(0.5, float(os.getenv("FILE_SEARCH_RETRY_BACKOFF_SEC", "1.5")))
        self.openai_timeout_sec = max(30.0, float(os.getenv("FILE_SEARCH_OPENAI_TIMEOUT_SEC", "180")))
        self.openai_client = OpenAI(
            api_key=self.openai_api_key,
            timeout=self.openai_timeout_sec,
            max_retries=2,
        )
        self.file_search_model = os.getenv("OPENAI_FILE_SEARCH_MODEL", "gpt-5.1")







    # ------------------------------------------------------------------



    # Metadata helpers



    # ------------------------------------------------------------------



    def _get_scope_dir(self, project_id: str, chat_id: Optional[str]) -> str:



        base = os.path.join(config.get_project_path(project_id), "file_search")



        if chat_id:



            return os.path.join(base, "chats", chat_id)



        return os.path.join(base, "project")







    def _get_metadata_path(self, project_id: str, chat_id: Optional[str]) -> str:
        return os.path.join(self._get_scope_dir(project_id, chat_id), "metadata.json")

    def _get_db_path(self) -> Optional[str]:
        if not config.user_data_path:
            return None
        db_path = os.path.join(config.user_data_path, DB_FILENAME)
        if os.path.exists(db_path):
            return db_path
        legacy_path = os.path.join(config.user_data_path, LEGACY_DB_FILENAME)
        if os.path.exists(legacy_path):
            return legacy_path
        return db_path

    def _connect_db(self) -> sqlite3.Connection:
        db_path = self._get_db_path()
        if not db_path:
            raise ValueError("user_data_path not set")
        conn = sqlite3.connect(db_path)
        conn.row_factory = sqlite3.Row
        return conn

    def _now_ms(self) -> int:
        return int(time.time() * 1000)

    def _to_ms(self, value: Any) -> Optional[int]:
        if value is None:
            return None
        if isinstance(value, (int, float)):
            return int(value)
        if isinstance(value, str):
            try:
                return int(value)
            except ValueError:
                pass
            try:
                parsed = datetime.fromisoformat(value)
                return int(parsed.timestamp() * 1000)
            except Exception:
                return None
        return None

    def _ms_to_iso(self, value: Optional[int]) -> Optional[str]:
        if value is None:
            return None
        try:
            return datetime.fromtimestamp(value / 1000).isoformat()
        except Exception:
            return None

    def _safe_json_loads(self, raw_json: Optional[str]) -> Dict[str, Any]:
        if not raw_json:
            return {}
        try:
            data = json.loads(raw_json)
            return data if isinstance(data, dict) else {}
        except Exception:
            return {}

    def _vector_store_row_id(self, project_id: str, chat_id: Optional[str]) -> str:
        scope = "chat" if chat_id else "project"
        suffix = chat_id if chat_id else "project"
        return f"{scope}:vector_store:{project_id}:{suffix}"

    def _file_row_id(self, project_id: str, chat_id: Optional[str], file_id: str) -> str:
        scope = "chat" if chat_id else "project"
        suffix = chat_id if chat_id else "project"
        return f"{scope}:file:{project_id}:{suffix}:{file_id}"

    def _delete_metadata_scope(self, project_id: str, chat_id: Optional[str]) -> None:
        db_path = self._get_db_path()
        if not db_path or not os.path.exists(db_path):
            return
        try:
            with self._connect_db() as conn:
                if chat_id:
                    conn.execute(
                        "DELETE FROM file_search_metadata WHERE project_id = ? AND chat_id = ?",
                        (project_id, chat_id),
                    )
                else:
                    conn.execute(
                        "DELETE FROM file_search_metadata WHERE project_id = ? AND chat_id IS NULL",
                        (project_id,),
                    )
                conn.commit()
        except Exception as e:
            debug(f"[FILE_SEARCH] Failed to delete metadata scope: {e}")

    def _load_metadata(self, project_id: str, chat_id: Optional[str]) -> Dict[str, Any]:
        db_path = self._get_db_path()
        if not db_path or not os.path.exists(db_path):
            debug("[FILE_SEARCH] DB not available; metadata empty")
            return {}

        try:
            with self._connect_db() as conn:
                if chat_id:
                    rows = conn.execute(
                        "SELECT * FROM file_search_metadata WHERE project_id = ? AND chat_id = ?",
                        (project_id, chat_id),
                    ).fetchall()
                else:
                    rows = conn.execute(
                        "SELECT * FROM file_search_metadata WHERE project_id = ? AND chat_id IS NULL",
                        (project_id,),
                    ).fetchall()
        except Exception as e:
            debug(f"[FILE_SEARCH] Failed to load metadata from DB: {e}")
            return {}

        metadata: Dict[str, Any] = {}
        for row in rows:
            entry = self._safe_json_loads(row["raw_json"])
            kind = row["kind"]
            if kind == "vector_store":
                entry.setdefault("vector_store_id", row["vector_store_id"])
                entry.setdefault("project_id", row["project_id"])
                entry.setdefault("chat_id", row["chat_id"])
                entry.setdefault("scope", row["scope"])
                if "updated_at" not in entry and row["updated_at"]:
                    entry["updated_at"] = self._ms_to_iso(row["updated_at"])
                metadata["__vector_store__"] = entry
                continue

            file_id = row["file_id"] or entry.get("file_id")
            if not file_id:
                continue
            entry.setdefault("file_id", file_id)
            entry.setdefault("file_name", row["file_name"])
            entry.setdefault("openai_file_id", row["openai_file_id"])
            entry.setdefault("vector_store_file_id", row["vector_store_file_id"])
            entry.setdefault("vector_store_id", row["vector_store_id"])
            entry.setdefault("usage_bytes", row["usage_bytes"])
            entry.setdefault("scope", row["scope"])
            if "uploaded_at" not in entry and row["uploaded_at"]:
                entry["uploaded_at"] = self._ms_to_iso(row["uploaded_at"])
            if "updated_at" not in entry and row["updated_at"]:
                entry["updated_at"] = self._ms_to_iso(row["updated_at"])

            metadata[self._file_key(file_id, chat_id)] = entry

        return metadata

    def _save_metadata(self, project_id: str, chat_id: Optional[str], metadata: Dict[str, Any]) -> None:
        db_path = self._get_db_path()
        if not db_path or not os.path.exists(db_path):
            debug("[FILE_SEARCH] DB not available; metadata save skipped")
            return

        scope_label = self._scope_label(chat_id)
        now_ms = self._now_ms()

        try:
            with self._connect_db() as conn:
                if chat_id:
                    conn.execute(
                        "DELETE FROM file_search_metadata WHERE project_id = ? AND chat_id = ?",
                        (project_id, chat_id),
                    )
                else:
                    conn.execute(
                        "DELETE FROM file_search_metadata WHERE project_id = ? AND chat_id IS NULL",
                        (project_id,),
                    )

                scope_info = metadata.get("__vector_store__")
                if isinstance(scope_info, dict):
                    updated_at_ms = self._to_ms(scope_info.get("updated_at")) or now_ms
                    conn.execute(
                        "INSERT INTO file_search_metadata (id, scope, project_id, chat_id, kind, file_id, file_name, openai_file_id, vector_store_file_id, vector_store_id, usage_bytes, uploaded_at, updated_at, raw_json) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
                        (
                            self._vector_store_row_id(project_id, chat_id),
                            scope_label,
                            project_id,
                            chat_id,
                            "vector_store",
                            None,
                            None,
                            None,
                            None,
                            scope_info.get("vector_store_id"),
                            None,
                            None,
                            updated_at_ms,
                            json.dumps(scope_info, ensure_ascii=False),
                        ),
                    )

                for key, entry in metadata.items():
                    if key == "__vector_store__":
                        continue
                    if not isinstance(entry, dict):
                        continue
                    file_id = entry.get("file_id")
                    if not file_id:
                        continue
                    uploaded_at_ms = self._to_ms(entry.get("uploaded_at"))
                    updated_at_ms = self._to_ms(entry.get("updated_at")) or uploaded_at_ms or now_ms
                    conn.execute(
                        "INSERT INTO file_search_metadata (id, scope, project_id, chat_id, kind, file_id, file_name, openai_file_id, vector_store_file_id, vector_store_id, usage_bytes, uploaded_at, updated_at, raw_json) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
                        (
                            self._file_row_id(project_id, chat_id, file_id),
                            entry.get("scope") or scope_label,
                            project_id,
                            chat_id,
                            "file",
                            file_id,
                            entry.get("file_name"),
                            entry.get("openai_file_id"),
                            entry.get("vector_store_file_id"),
                            entry.get("vector_store_id"),
                            entry.get("usage_bytes"),
                            uploaded_at_ms,
                            updated_at_ms,
                            json.dumps(entry, ensure_ascii=False),
                        ),
                    )
                conn.commit()
        except Exception as e:
            debug(f"[FILE_SEARCH] Failed to save metadata to DB: {e}")

    def _file_key(self, file_id: str, chat_id: Optional[str]) -> str:



        scope = f"chat_file:{chat_id}" if chat_id else "project_file"



        return f"{scope}:{file_id}"







    def _scope_label(self, chat_id: Optional[str]) -> str:



        return "chat" if chat_id else "project"

    def _ensure_chat_vector_store_expiration(
        self,
        project_id: str,
        chat_id: Optional[str],
        vector_store_id: str,
        metadata: Dict[str, Any],
    ) -> None:
        if not chat_id or CHAT_VECTOR_STORE_TTL_DAYS <= 0:
            return
        scope_info = metadata.get("__vector_store__", {})
        if scope_info.get("expires_after_days") == CHAT_VECTOR_STORE_TTL_DAYS:
            return

        client = self.openai_client
        if not client:
            return

        try:
            client.vector_stores.update(
                vector_store_id=vector_store_id,
                expires_after={
                    "anchor": "last_active_at",
                    "days": CHAT_VECTOR_STORE_TTL_DAYS
                }
            )
            scope_info["expires_after_days"] = CHAT_VECTOR_STORE_TTL_DAYS
            metadata["__vector_store__"] = scope_info
            self._save_metadata(project_id, chat_id, metadata)
            debug(f"[FILE_SEARCH] Chat vector store 만료 설정 적용: vs_id={vector_store_id}, days={CHAT_VECTOR_STORE_TTL_DAYS}")
        except Exception as e:
            debug(f"[FILE_SEARCH] Chat vector store 만료 설정 실패: {e}")

    def _has_files(self, metadata: Dict[str, Any]) -> bool:



        for key in metadata.keys():



            if key.startswith("project_file:") or key.startswith("chat_file:"):



                return True



        return False







    # ------------------------------------------------------------------



    # Vector store helpers



    # ------------------------------------------------------------------



    def _ensure_openai(self):



        if not self.openai_client:



            raise ValueError("OPENAI_API_KEY_REQUIRED")



        return self.openai_client

    def _is_retryable_error(self, err: Exception) -> bool:
        name = err.__class__.__name__.lower()
        message = str(err).lower()
        retryable_tokens = (
            "timeout",
            "connect",
            "connection",
            "temporar",
            "503",
            "502",
            "504",
            "rate limit",
            "ratelimit",
            "apiconnection",
            "apitimeout",
            "internalservererror",
        )
        return any(token in name or token in message for token in retryable_tokens)

    def _run_with_retry(self, label: str, fn: Callable[[], Any]) -> Any:
        last_error: Optional[Exception] = None
        for attempt in range(1, self.max_network_retries + 1):
            try:
                return fn()
            except Exception as err:
                last_error = err
                is_retryable = self._is_retryable_error(err)
                if not is_retryable or attempt >= self.max_network_retries:
                    raise
                wait_sec = self.retry_backoff_sec * attempt
                debug(
                    f"[FILE_SEARCH] {label} 실패 ({attempt}/{self.max_network_retries}) - "
                    f"{wait_sec:.1f}초 후 재시도: {err}"
                )
                time.sleep(wait_sec)
        if last_error is not None:
            raise last_error
        raise RuntimeError(f"{label} 재시도 실패")







    def _get_or_create_vector_store(self, project_id: str, chat_id: Optional[str]) -> Tuple[str, Dict[str, Any]]:



        metadata = self._load_metadata(project_id, chat_id)



        scope_info = metadata.get("__vector_store__")



        if scope_info and scope_info.get("vector_store_id"):



            vector_store_id = scope_info["vector_store_id"]
            self._ensure_chat_vector_store_expiration(project_id, chat_id, vector_store_id, metadata)
            return vector_store_id, metadata







        client = self._ensure_openai()



        scope_label = self._scope_label(chat_id)



        vs_name = f"{VECTOR_STORE_PREFIX}-{scope_label}-{project_id}" + (f"-{chat_id}" if chat_id else "")



        vector_store = self._run_with_retry(
            "vector_stores.create",
            lambda: client.vector_stores.create(
                name=vs_name,
                metadata={
                    "project_id": project_id,
                    "chat_id": chat_id or "",
                    "scope": scope_label,
                },
            ),
        )



        metadata["__vector_store__"] = {



            "vector_store_id": vector_store.id,



            "project_id": project_id,



            "chat_id": chat_id,



            "scope": scope_label,



            "updated_at": datetime.now().isoformat(),



        }



        self._ensure_chat_vector_store_expiration(project_id, chat_id, vector_store.id, metadata)
        self._save_metadata(project_id, chat_id, metadata)



        return vector_store.id, metadata







    def _get_scope_metadata(self, project_id: str, chat_id: Optional[str]) -> Tuple[Optional[str], Dict[str, Any]]:



        metadata = self._load_metadata(project_id, chat_id)



        scope_info = metadata.get("__vector_store__", {})



        vector_store_id = scope_info.get("vector_store_id")



        if vector_store_id and chat_id:
            self._ensure_chat_vector_store_expiration(project_id, chat_id, vector_store_id, metadata)
        return vector_store_id, metadata







    def _find_file_entry_by_name(



        self, metadata: Dict[str, Any], chat_id: Optional[str], file_name: str



    ) -> Optional[Dict[str, Any]]:



        prefix = f"chat_file:{chat_id}:" if chat_id else "project_file:"



        for key, entry in metadata.items():



            if not key.startswith(prefix):



                continue



            if isinstance(entry, dict) and entry.get("file_name") == file_name:



                return entry



        return None







    def get_vector_store_ids(self, project_id: str, chat_id: Optional[str] = None) -> Tuple[List[str], bool]:



        ids: List[str] = []



        has_files = False







        project_meta = self._load_metadata(project_id, None)



        if project_meta:



            scope_info = project_meta.get("__vector_store__", {})



            vs_id = scope_info.get("vector_store_id")



            if vs_id:



                ids.append(vs_id)



            if self._has_files(project_meta):



                has_files = True







        if chat_id:



            chat_meta = self._load_metadata(project_id, chat_id)



            if chat_meta:



                scope_info = chat_meta.get("__vector_store__", {})



                vs_id = scope_info.get("vector_store_id")



                if vs_id:



                    ids.append(vs_id)



                if self._has_files(chat_meta):



                    has_files = True







        return ids, has_files







    # ------------------------------------------------------------------



    # Indexing / Deletion



    # ------------------------------------------------------------------



    def index_file(



        self,



        project_id: str,



        file_id: str,



        text_content: str,



        file_name: str,



        chat_id: Optional[str] = None,



        progress_callback: Optional[Callable[[float, str], None]] = None,



        file_path: Optional[str] = None,



    ) -> Dict[str, Any]:



        if not project_id or not file_id:



            return {"success": False, "error": "Missing required parameters (project_id, file_id)"}







        if not (text_content or file_path):



            return {"success": False, "error": "Missing text content or file path"}
        if file_path:
            ext = os.path.splitext(file_path)[1].lower()

            if ext in UNSUPPORTED_IMAGE_EXTS:
                return {"success": False, "error": "Image files are not supported for indexing"}

        client = self._ensure_openai()



        scope_label = self._scope_label(chat_id)



        debug(f"[FILE_SEARCH] 인덱싱 시작: {file_name} (scope={scope_label}, file_id={file_id})")







        if progress_callback:



            progress_callback(0.05, "파일 업로드 준비 중...")







        vector_store_id, metadata = self._get_or_create_vector_store(project_id, chat_id)



        debug(f"[FILE_SEARCH] Vector Store 확보: {vector_store_id}")







        # 기존 항목 삭제



        existing_key = self._file_key(file_id, chat_id)



        if existing_key in metadata:



            self.delete_by_file_id(project_id, file_id, chat_id)



            metadata = self._load_metadata(project_id, chat_id)







        upload_bytes, upload_filename, attributes = self._build_upload_payload(



            text_content=text_content,



            file_name=file_name,



            file_id=file_id,



            chat_id=chat_id,



            file_path=file_path,



        )







        if progress_callback:



            progress_callback(0.2, "파일 업로드 중...")







        debug(f"[FILE_SEARCH] OpenAI Files API 호출: {upload_filename} ({len(upload_bytes):,} bytes)")



        uploaded_file = self._run_with_retry(
            "files.create",
            lambda: client.files.create(
                file=(upload_filename, upload_bytes),
                purpose="assistants",
            ),
        )



        debug(f"[FILE_SEARCH] 파일 업로드 완료: openai_file_id={uploaded_file.id}")







        if progress_callback:



            progress_callback(0.4, "파일 인덱싱 중...")







        debug(f"[FILE_SEARCH] Vector Store에 파일 추가 중...")



        vector_file = self._run_with_retry(
            "vector_stores.files.create",
            lambda: client.vector_stores.files.create(
                vector_store_id=vector_store_id,
                file_id=uploaded_file.id,
                attributes=attributes,
            ),
        )



        debug(f"[FILE_SEARCH] Vector Store 파일 생성: vs_file_id={vector_file.id}, status={vector_file.status}")







        vector_file = self._poll_vector_store_file(



            vector_store_id=vector_store_id,



            vector_store_file_id=vector_file.id,



            progress_callback=progress_callback,



        )







        if vector_file.status != "completed":



            return {"success": False, "error": f"Vector store file status: {vector_file.status}"}







        # 스토리지 비용 로그



        log_file_search_storage(



            usage_bytes=vector_file.usage_bytes,



            operation="file_search:upload",



            context_info=f"{file_name} ({self._scope_label(chat_id)})",



        )







        metadata[self._file_key(file_id, chat_id)] = {



            "file_id": file_id,



            "file_name": file_name,



            "openai_file_id": uploaded_file.id,



            "vector_store_file_id": vector_file.id,



            "vector_store_id": vector_store_id,



            "usage_bytes": vector_file.usage_bytes,



            "scope": self._scope_label(chat_id),



            "uploaded_at": datetime.now().isoformat(),



        }



        self._save_metadata(project_id, chat_id, metadata)







        if progress_callback:



            progress_callback(1.0, "인덱싱 완료")







        debug(f"[FILE_SEARCH] ✅ 인덱싱 성공: {file_name} → vs_id={vector_store_id}, usage={vector_file.usage_bytes} bytes")



        # 비용 계산 (일일 비용)
        COST_PER_GB_DAY = 0.10
        cost_per_day = (vector_file.usage_bytes / 1_073_741_824) * COST_PER_GB_DAY

        return {
            "success": True,
            "indexed_count": 1,
            "vector_store_id": vector_store_id,
            "usage": {
                "bytes": vector_file.usage_bytes,
                "cost_usd_per_day": cost_per_day
            }
        }







    def delete_by_file_id(self, project_id: str, file_id: str, chat_id: Optional[str] = None) -> Dict[str, Any]:



        scope_label = self._scope_label(chat_id)



        debug(f"[FILE_SEARCH] 삭제 시작: file_id={file_id}, scope={scope_label}")



        



        metadata = self._load_metadata(project_id, chat_id)



        key = self._file_key(file_id, chat_id)



        entry = metadata.get(key)



        if not entry:



            debug(f"[FILE_SEARCH] ❌ 메타데이터에서 파일을 찾을 수 없음: {file_id}")



            return {"success": False, "error": f"File {file_id} not found in metadata"}







        client = self.openai_client



        if not client:



            return {"success": False, "error": "OPENAI_API_KEY_REQUIRED"}



        vector_store_id = entry.get("vector_store_id")



        vector_store_file_id = entry.get("vector_store_file_id")



        openai_file_id = entry.get("openai_file_id")







        warnings: List[str] = []



        vector_deleted = False



        file_deleted = False







        try:



            if vector_store_id and vector_store_file_id:



                debug(f"[FILE_SEARCH] OpenAI Vector Store 파일 삭제 호출: vs_id={vector_store_id}, vs_file_id={vector_store_file_id}")



                client.vector_stores.files.delete(vector_store_id=vector_store_id, file_id=vector_store_file_id)



                vector_deleted = True



                debug(f"[FILE_SEARCH] ✅ Vector Store 파일 삭제 완료")



            else:



                return {"success": False, "error": "vector_store_id or vector_store_file_id missing"}



        except Exception as e:



            debug(f"[FILE_SEARCH] ❌ Vector store file delete failed: {e}")



            return {"success": False, "error": f"vector_store delete failed: {e}"}







        try:



            if openai_file_id:



                debug(f"[FILE_SEARCH] OpenAI Files API 삭제 호출: file_id={openai_file_id}")



                client.files.delete(openai_file_id)



                file_deleted = True



                debug(f"[FILE_SEARCH] ✅ OpenAI 파일 삭제 완료: {openai_file_id}")



            else:



                warnings.append("openai_file_id missing")



        except Exception as e:



            debug(f"[FILE_SEARCH] File delete failed: {e}")



            return {



                "success": False,



                "error": f"openai_file delete failed: {e}",



                "vector_store_deleted": vector_deleted,



                "openai_file_deleted": False,



                "metadata_removed": False,



                "warnings": warnings



            }







        del metadata[key]



        self._save_metadata(project_id, chat_id, metadata)



        



        file_name = entry.get("file_name", file_id)



        debug(f"[FILE_SEARCH] ✅ 삭제 완료: {file_name} (vector={vector_deleted}, file={file_deleted})")



        return {



            "success": vector_deleted and (file_deleted or not openai_file_id),



            "vector_store_deleted": vector_deleted,



            "openai_file_deleted": file_deleted,



            "metadata_removed": True,



            "warnings": warnings



        }







    def delete_chat_scope(self, project_id: str, chat_id: str) -> Dict[str, Any]:
        metadata = self._load_metadata(project_id, chat_id)
        client = self.openai_client

        vector_store_id = metadata.get("__vector_store__", {}).get("vector_store_id")

        for key, entry in list(metadata.items()):
            if not key.startswith(f"chat_file:{chat_id}:"):
                continue

            openai_file_id = entry.get("openai_file_id")
            vector_store_file_id = entry.get("vector_store_file_id")

            try:
                if client and vector_store_id and vector_store_file_id:
                    client.vector_stores.files.delete(
                        vector_store_id=vector_store_id,
                        file_id=vector_store_file_id,
                    )
            except Exception as e:
                debug(f"[FILE_SEARCH] Vector store file delete failed: {e}")

            try:
                if client and openai_file_id:
                    client.files.delete(openai_file_id)
            except Exception as e:
                debug(f"[FILE_SEARCH] File delete failed: {e}")

        try:
            if client and vector_store_id:
                client.vector_stores.delete(vector_store_id)
        except Exception as e:
            debug(f"[FILE_SEARCH] Vector store delete failed: {e}")

        self._delete_metadata_scope(project_id, chat_id)

        metadata_path = self._get_metadata_path(project_id, chat_id)
        if os.path.exists(metadata_path):
            try:
                os.remove(metadata_path)
            except Exception as e:
                debug(f"[FILE_SEARCH] Metadata delete failed: {e}")

        return {"success": True, "deleted_path": self._get_db_path() or metadata_path}

    def delete_project_scope(self, project_id: str) -> Dict[str, Any]:
        metadata = self._load_metadata(project_id, None)
        client = self.openai_client
        warnings: List[str] = []

        if not client:
            warnings.append("OPENAI_API_KEY_REQUIRED")

        vector_store_id = metadata.get("__vector_store__", {}).get("vector_store_id")
        vector_store_files_deleted = 0
        openai_files_deleted = 0
        project_files_seen = 0

        for key, entry in list(metadata.items()):
            if not key.startswith("project_file:"):
                continue
            project_files_seen += 1

            vector_store_file_id = entry.get("vector_store_file_id")
            openai_file_id = entry.get("openai_file_id")

            if client and vector_store_id and vector_store_file_id:
                try:
                    client.vector_stores.files.delete(
                        vector_store_id=vector_store_id,
                        file_id=vector_store_file_id,
                    )
                    vector_store_files_deleted += 1
                except Exception as e:
                    warnings.append(f"Vector store file delete failed: {e}")
            elif client and vector_store_id and not vector_store_file_id:
                warnings.append("vector_store_file_id missing for project file")

            if client and openai_file_id:
                try:
                    client.files.delete(openai_file_id)
                    openai_files_deleted += 1
                except Exception as e:
                    warnings.append(f"OpenAI file delete failed: {e}")
            elif client and not openai_file_id:
                warnings.append("openai_file_id missing for project file")

        project_vector_store_deleted = False
        if client and vector_store_id:
            try:
                client.vector_stores.delete(vector_store_id)
                project_vector_store_deleted = True
            except Exception as e:
                warnings.append(f"Vector store delete failed: {e}")

        self._delete_metadata_scope(project_id, None)

        chat_scopes_deleted = 0
        chat_ids: List[str] = []
        db_path = self._get_db_path()
        if db_path and os.path.exists(db_path):
            try:
                with self._connect_db() as conn:
                    rows = conn.execute(
                        "SELECT DISTINCT chat_id FROM file_search_metadata WHERE project_id = ? AND chat_id IS NOT NULL",
                        (project_id,),
                    ).fetchall()
                chat_ids = [row["chat_id"] for row in rows if row["chat_id"]]
            except Exception as e:
                warnings.append(f"Chat scope list failed: {e}")

        for chat_id in chat_ids:
            try:
                result = self.delete_chat_scope(project_id, chat_id)
                if result.get("success"):
                    chat_scopes_deleted += 1
                else:
                    warnings.append(f"Chat scope delete failed: {chat_id}")
            except Exception as e:
                warnings.append(f"Chat scope delete failed ({chat_id}): {e}")

        return {
            "success": True,
            "project_files_seen": project_files_seen,
            "vector_store_files_deleted": vector_store_files_deleted,
            "openai_files_deleted": openai_files_deleted,
            "project_vector_store_deleted": project_vector_store_deleted,
            "chat_scopes_deleted": chat_scopes_deleted,
            "deleted_path": db_path,
            "warnings": warnings,
        }

    def delete_by_pair_id(self, project_id: str, pair_id: str, chat_id: Optional[str] = None) -> Dict[str, Any]:



        # Template pair indexing is disabled (no-op)



        return {"success": True, "message": "Template pair indexing is disabled"}







    # ------------------------------------------------------------------



    # Optional search helper (for backward compatibility)



    # ------------------------------------------------------------------



    def query_multi_scope(



        self,



        query_text: str,



        project_id: str,



        chat_id: Optional[str] = None,



        file_name: Optional[str] = None,



        top_k: int = 20,



        mode: str = "file_search",



        progress_callback: Optional[Callable[[str, str], None]] = None,



    ) -> Dict[str, Any]:



        if progress_callback:



            if file_name:



                progress_callback("search", f"'{file_name}' 파일 검색 중...")



            else:



                progress_callback("search", "업로드 자료 검색 중...")







        if not self.openai_client:



            return {"success": False, "error": "OPENAI_API_KEY_REQUIRED", "results": []}







        results: List[Dict[str, Any]] = []



        found_sources: List[str] = []







        scopes_to_search: List[Tuple[str, str]] = []







        def add_scope(scope_chat_id: Optional[str], scope_label: str) -> None:



            vector_store_id, metadata = self._get_scope_metadata(project_id, scope_chat_id)



            if not vector_store_id or not self._has_files(metadata):



                return



            if file_name and not self._find_file_entry_by_name(metadata, scope_chat_id, file_name):



                return



            scopes_to_search.append((vector_store_id, scope_label))







        if file_name:



            if chat_id:



                add_scope(chat_id, "chat_file")



            if not scopes_to_search:



                add_scope(None, "project_file")



        else:



            if chat_id:



                add_scope(chat_id, "chat_file")



            add_scope(None, "project_file")







        if not scopes_to_search:



            if progress_callback:



                progress_callback("complete", "검색 완료 (관련 자료 없음)")



            return {"success": True, "results": [], "mode": mode}







        filters = None



        if file_name:



            filters = {"type": "eq", "key": "source_name", "value": file_name}







        for vector_store_id, scope_label in scopes_to_search:



            scope_hits = self._search_vector_store(



                vector_store_id=vector_store_id,



                query_text=query_text,



                top_k=top_k,



                filters=filters,



            )



            for hit in scope_hits:



                source_name = hit.get("attributes", {}).get("source_name") or hit.get("filename")



                if not source_name:



                    source_name = file_name or "unknown"



                if file_name and source_name != file_name:



                    continue



                content_text = (hit.get("text") or "")[:8000]



                if not content_text.strip():



                    continue



                results.append({



                    "scope": scope_label,



                    "content": content_text,



                    "source": source_name,



                })



                if source_name not in found_sources:



                    found_sources.append(source_name)







        if progress_callback:



            if found_sources:



                sources_str = ", ".join(found_sources[:3])



                if len(found_sources) > 3:



                    sources_str += f" 외 {len(found_sources) - 3}개"



                progress_callback("complete", f"자료 검색 완료: {sources_str}")



            else:



                progress_callback("complete", "검색 완료 (관련 자료 없음)")







        return {"success": True, "results": results, "mode": mode}







    def _search_vector_store(



        self,



        vector_store_id: str,



        query_text: str,



        top_k: int,



        filters: Optional[Dict[str, Any]] = None,



    ) -> List[Dict[str, Any]]:



        response_results = self._search_with_responses(



            vector_store_id=vector_store_id,



            query_text=query_text,



            top_k=top_k,



            filters=filters,



        )



        if response_results is None:



            debug("[FILE_SEARCH] Responses API search returned no file_search_call; skipping fallback")



            return []



        return response_results







    def _search_with_responses(



        self,



        vector_store_id: str,



        query_text: str,



        top_k: int,



        filters: Optional[Dict[str, Any]] = None,



    ) -> Optional[List[Dict[str, Any]]]:



        client = self._ensure_openai()



        try:



            tool_params: Dict[str, Any] = {



                "type": "file_search",



                "vector_store_ids": [vector_store_id],



                "max_num_results": top_k,



            }



            if filters:



                tool_params["filters"] = filters







            response = client.responses.create(



                model=self.file_search_model,



                input=query_text,



                tools=[tool_params],



                tool_choice={"type": "file_search"},



                include=["file_search_call.results"],



                max_output_tokens=8000,



            )



        except Exception as e:



            debug(f"[FILE_SEARCH] Responses API search failed: {e}")



            return None







        usage = getattr(response, "usage", None)



        input_tokens = None



        output_tokens = None



        total_tokens = None



        if usage:



            if isinstance(usage, dict):



                input_tokens = usage.get("input_tokens")



                output_tokens = usage.get("output_tokens")



                total_tokens = usage.get("total_tokens")



            else:



                input_tokens = getattr(usage, "input_tokens", None)



                output_tokens = getattr(usage, "output_tokens", None)



                total_tokens = getattr(usage, "total_tokens", None)







        log_file_search_tool_usage(



            tool_calls=1,



            operation="file_search:query",



            context_info=(



                f"vector_store_id={vector_store_id}, "



                f"top_k={top_k}, query_len={len(query_text) if query_text else 0}"



            ),



            input_tokens=input_tokens,



            output_tokens=output_tokens,



            total_tokens=total_tokens,



        )







        results: List[Dict[str, Any]] = []



        saw_call = False



        for item in getattr(response, "output", []) or []:



            if getattr(item, "type", "") != "file_search_call":



                continue



            saw_call = True



            for res in item.results or []:



                results.append({



                    "filename": res.filename,



                    "text": res.text or "",



                    "attributes": res.attributes or {},



                })







        if not saw_call:



            debug("[FILE_SEARCH] Responses API output missing file_search_call")



            return None



        return results







    # ------------------------------------------------------------------



    # Internal helpers



    # ------------------------------------------------------------------



    def _convert_hwp_to_docx(self, hwp_path: str) -> Optional[bytes]:



        """HWP → DOCX 변환 (pyhwpx 사용)"""



        try:



            import pyhwpx



            import tempfile







            hwp = pyhwpx.Hwp()



            hwp.Open(hwp_path)







            # 임시 DOCX 파일 생성



            with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as tmp:



                tmp_docx = tmp.name







            # DOCX로 저장



            hwp.SaveAs(tmp_docx, format="DOCX")



            hwp.Quit()







            # 바이너리 읽기



            with open(tmp_docx, "rb") as f:



                docx_bytes = f.read()







            # 임시 파일 삭제



            os.remove(tmp_docx)







            return docx_bytes



        except Exception as e:



            debug(f"[FILE_SEARCH] HWP → DOCX 변환 실패: {e}")



            return None







    def _convert_excel_to_markdown(self, excel_path: str) -> Optional[str]:



        """Excel → Markdown 표 변환 (openpyxl 사용, 병합 셀 지원)"""



        try:



            from openpyxl import load_workbook







            workbook = load_workbook(excel_path, data_only=True)



            markdown_content = []







            for sheet_name in workbook.sheetnames:



                sheet = workbook[sheet_name]







                if sheet.max_row == 0:



                    continue







                markdown_content.append(f"## {sheet_name}\n")







                # 병합 셀 정보 저장



                merged_cells = {}



                for merged_range in sheet.merged_cells.ranges:



                    min_row, min_col = merged_range.min_row, merged_range.min_col



                    for row in range(merged_range.min_row, merged_range.max_row + 1):



                        for col in range(merged_range.min_col, merged_range.max_col + 1):



                            merged_cells[(row, col)] = (min_row, min_col)







                # 표 데이터 추출



                rows_data = []



                for row_idx, row in enumerate(sheet.iter_rows(min_row=1, max_row=sheet.max_row), start=1):



                    row_values = []



                    for col_idx, cell in enumerate(row, start=1):



                        if (row_idx, col_idx) in merged_cells:



                            # 병합 셀: 원본 셀 값 사용



                            origin_row, origin_col = merged_cells[(row_idx, col_idx)]



                            origin_cell = sheet.cell(origin_row, origin_col)



                            value = origin_cell.value if origin_cell.value is not None else ""



                        else:



                            value = cell.value if cell.value is not None else ""



                        row_values.append(str(value))



                    rows_data.append(row_values)







                if not rows_data:



                    continue







                # Markdown 표 생성



                max_cols = max(len(row) for row in rows_data)



                for row in rows_data:



                    while len(row) < max_cols:



                        row.append("")







                # 헤더 행



                markdown_content.append("| " + " | ".join(rows_data[0]) + " |")



                markdown_content.append("| " + " | ".join(["---"] * max_cols) + " |")







                # 데이터 행



                for row in rows_data[1:]:



                    markdown_content.append("| " + " | ".join(row) + " |")







                markdown_content.append("\n")







            return "\n".join(markdown_content)



        except Exception as e:



            debug(f"[FILE_SEARCH] Excel → Markdown 변환 실패: {e}")



            return None







    def _build_upload_payload(



        self,



        text_content: str,



        file_name: str,



        file_id: str,



        chat_id: Optional[str],



        file_path: Optional[str],



    ) -> Tuple[bytes, str, Dict[str, Any]]:



        base_name = os.path.splitext(file_name)[0] if file_name else f"file_{file_id}"



        attributes: Dict[str, Any] = {



            "source_name": file_name,  # 원본 파일명 유지



            "file_id": file_id,



            "scope": self._scope_label(chat_id),



        }







        if file_path and os.path.exists(file_path):



            ext = os.path.splitext(file_path)[1].lower()







            # 1. HWP/HWPX는 별도 추출(text_content) 결과를 우선 사용
            #    (재오픈/종료 과정에서 바인딩된 한글 문서가 닫히는 부작용 방지)
            if ext in CONVERT_TO_DOCX:
                upload_filename = f"{base_name}{TEXT_FALLBACK_EXT}"
                cleaned_text = (text_content or "").strip()

                # 프론트엔드 FileReader.readAsText()로 HWP 바이너리를 읽으면
                # 빈/깨진 텍스트가 됨 → 인덱싱 성공해도 검색 결과 0건
                # 최소 길이 가드 + pyhwpx 변환 fallback
                MIN_HWP_TEXT_LEN = 64
                if len(cleaned_text) < MIN_HWP_TEXT_LEN:
                    debug(
                        f"[FILE_SEARCH] HWP text_content 짧음 "
                        f"(len={len(cleaned_text)} < {MIN_HWP_TEXT_LEN}) "
                        f"- pyhwpx 변환 시도: {file_name}"
                    )
                    try:
                        docx_bytes = self._convert_hwp_to_docx(file_path)
                    except Exception as exc:
                        debug(f"[FILE_SEARCH] HWP→DOCX 변환 실패: {exc}")
                        docx_bytes = None

                    if docx_bytes:
                        attributes["converted_from"] = ext
                        attributes["conversion_mode"] = "pyhwpx_docx"
                        return docx_bytes, f"{base_name}.docx", attributes

                    debug(
                        f"[FILE_SEARCH] HWP 텍스트 추출 실패 - 검색 품질 저하 가능: {file_name}"
                    )

                attributes["converted_from"] = ext
                attributes["conversion_mode"] = "text_fallback"
                return cleaned_text.encode("utf-8"), upload_filename, attributes







            # 2. Native 지원 확장자: 원본 그대로 업로드



            if ext in FILE_SEARCH_NATIVE_EXTS:



                with open(file_path, "rb") as f:



                    return f.read(), os.path.basename(file_path), attributes







            # 3. Excel → Markdown 표 변환



            if ext in CONVERT_TO_MD:



                markdown_content = self._convert_excel_to_markdown(file_path)



                if markdown_content:



                    upload_filename = f"{base_name}.md"



                    attributes["converted_from"] = ext  # 변환 정보 저장



                    return markdown_content.encode("utf-8"), upload_filename, attributes







        # Fallback: 텍스트 컨텐츠



        upload_filename = f"{base_name}{TEXT_FALLBACK_EXT}"



        content = text_content or ""



        return content.encode("utf-8"), upload_filename, attributes







    def _poll_vector_store_file(



        self,



        vector_store_id: str,



        vector_store_file_id: str,



        progress_callback: Optional[Callable[[float, str], None]] = None,



    ):



        client = self._ensure_openai()



        start_time = time.time()



        last_progress = 0.4







        while True:



            vector_file = self._run_with_retry(
                "vector_stores.files.retrieve",
                lambda: client.vector_stores.files.retrieve(
                    vector_store_id=vector_store_id,
                    file_id=vector_store_file_id,
                ),
            )







            if vector_file.status in ("completed", "failed", "cancelled"):



                return vector_file







            elapsed = time.time() - start_time



            if elapsed > POLL_TIMEOUT_SEC:



                return vector_file







            if progress_callback:



                progress = min(0.9, last_progress + (elapsed / POLL_TIMEOUT_SEC) * 0.5)



                progress_callback(progress, "인덱싱 진행 중...")







            time.sleep(POLL_INTERVAL_SEC)











# Backward compatibility alias



FileSearchServiceV2 = FileSearchService



# Legacy alias for backward compatibility



RAGService = FileSearchService



RAGServiceV2 = FileSearchService



