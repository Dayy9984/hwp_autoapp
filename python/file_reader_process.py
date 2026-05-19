"""
File Reader Process - 별도 프로세스에서 파일 읽기 수행

hwp_com_process.py의 메인 RPC 루프를 blocking하지 않도록
파일 읽기/CVD 추출을 독립 프로세스에서 처리합니다.

Protocol: JSON-RPC over stdin/stdout (hwp_com_process.py와 동일)
"""

import json
import os
import sys
import traceback


def handle_request(request: dict) -> dict:
    method = request.get("method", "")
    params = request.get("params", {})
    request_id = request.get("id")

    result = {"id": request_id}

    try:
        if method == "ping":
            result["result"] = {"pong": True}

        elif method == "readTxtFile":
            from readers.txt_reader import TxtReader
            result["result"] = TxtReader.read(
                params.get("filePath", ""),
                params.get("maxChars")
            )

        elif method == "readExcelFile":
            from readers.excel_reader import ExcelReader
            result["result"] = ExcelReader.read(
                params.get("filePath", ""),
                params.get("sheetName")
            )

        elif method == "getExcelSheets":
            from readers.excel_reader import ExcelReader
            result["result"] = ExcelReader.get_sheet_names(
                params.get("filePath", "")
            )

        elif method == "readPdfFile":
            try:
                from services.cvd_service import CVDService
                cvd_service = CVDService(
                    lambda level, msg: print(f"[CVD][{level}] {msg}", file=sys.stderr)
                )
                pdf_result = cvd_service.extract_pdf(params.get("filePath", ""))
                if pdf_result.get("success"):
                    result["result"] = {
                        "success": True,
                        "text": pdf_result.get("text", ""),
                        "page_count": pdf_result.get("page_count", 0)
                    }
                else:
                    result["result"] = {
                        "success": False,
                        "error": pdf_result.get("error", "Unknown error")
                    }
            except Exception as e:
                result["result"] = {"success": False, "error": str(e)}

        elif method == "readHwpFile":
            try:
                file_path = params.get("filePath", "")
                ext = os.path.splitext(file_path)[1].lower()

                if ext in ['.hwp', '.hwpx']:
                    from services.cvd_service import CVDService
                    import tempfile
                    import shutil

                    cvd_service = CVDService(
                        lambda level, msg: print(f"[CVD][{level}] {msg}", file=sys.stderr),
                        allow_existing_instance=False
                    )

                    def progress_callback(pct: float, message: str):
                        progress_event = {
                            "type": "progress",
                            "event": "file:upload:progress",
                            "data": {"progress": pct, "message": message, "filePath": file_path}
                        }
                        print(json.dumps(progress_event, ensure_ascii=False), flush=True)

                    temp_dir = tempfile.mkdtemp(prefix="rag_hwp_")
                    try:
                        cvd_result = cvd_service.extract_single_cvd(
                            file_path, temp_dir, progress_callback=progress_callback
                        )
                        if cvd_result and cvd_result.get("success"):
                            cvd_path = cvd_result.get("cvd_path")
                            if cvd_path and os.path.exists(cvd_path):
                                with open(cvd_path, "r", encoding="utf-8") as f:
                                    text_content = f.read()
                                result["result"] = {"success": True, "text": text_content}
                            else:
                                result["result"] = {"success": False, "error": "CVD output missing"}
                        else:
                            error_message = cvd_result.get("error") if cvd_result else "HWP extraction failed"
                            result["result"] = {"success": False, "error": error_message}
                    finally:
                        shutil.rmtree(temp_dir, ignore_errors=True)
                else:
                    result["result"] = {"success": False, "error": f"Unsupported extension: {ext}"}

            except Exception as e:
                result["result"] = {"success": False, "error": str(e)}

        elif method == "readDocFile":
            try:
                file_path = params.get("filePath", "")
                ext = os.path.splitext(file_path)[1].lower()

                if ext == '.docx':
                    try:
                        from docx import Document
                        doc = Document(file_path)
                        texts = [para.text for para in doc.paragraphs if para.text.strip()]
                        text_content = '\n'.join(texts)
                        result["result"] = {"success": True, "text": text_content}
                    except ImportError:
                        import zipfile
                        import xml.etree.ElementTree as ET

                        texts = []
                        with zipfile.ZipFile(file_path, 'r') as zf:
                            if 'word/document.xml' in zf.namelist():
                                with zf.open('word/document.xml') as f:
                                    tree = ET.parse(f)
                                    root = tree.getroot()
                                    for t in root.iter(
                                        '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t'
                                    ):
                                        if t.text:
                                            texts.append(t.text)

                        text_content = ''.join(texts)
                        result["result"] = {"success": True, "text": text_content}
                elif ext == '.doc':
                    result["result"] = {
                        "success": False,
                        "error": "Legacy .doc format is not supported. Please convert to .docx."
                    }
                else:
                    result["result"] = {"success": False, "error": f"Unsupported extension: {ext}"}

            except Exception as e:
                result["result"] = {"success": False, "error": str(e)}

        elif method == "readPptFile":
            try:
                file_path = params.get("filePath", "")
                ext = os.path.splitext(file_path)[1].lower()

                if ext == '.pptx':
                    try:
                        from pptx import Presentation
                    except ImportError:
                        result["result"] = {"success": False, "error": "python-pptx not installed"}
                        return result

                    prs = Presentation(file_path)
                    parts = []
                    for i, slide in enumerate(prs.slides, 1):
                        parts.append(f"## Slide {i}")
                        for shape in slide.shapes:
                            if hasattr(shape, "text") and shape.text:
                                parts.append(shape.text)

                    text_content = "\n".join(parts)
                    result["result"] = {"success": True, "text": text_content}
                elif ext == '.ppt':
                    result["result"] = {
                        "success": False,
                        "error": "Legacy .ppt format is not supported. Please convert to .pptx."
                    }
                else:
                    result["result"] = {"success": False, "error": f"Unsupported extension: {ext}"}
            except Exception as e:
                result["result"] = {"success": False, "error": str(e)}

        elif method == "cvd:extractPair":
            try:
                from services.cvd_service import CVDService
                from services.config import config

                project_id = params.get("projectId")
                pair_id = params.get("pairId")
                template_path = params.get("templatePath")
                filled_path = params.get("filledPath")
                user_data_path = params.get("userDataPath")

                if not all([project_id, pair_id, template_path, filled_path, user_data_path]):
                    result["result"] = {"success": False, "error": "Missing required parameters"}
                else:
                    config.set_user_data_path(user_data_path)

                    def log_callback(level: str, message: str):
                        print(f"[CVDService][{level}] {message}", file=sys.stderr)

                    def progress_callback(progress: float, message: str):
                        event = {
                            "type": "progress",
                            "event": "cvd:progress",
                            "data": {
                                "pairId": pair_id,
                                "progress": int(progress * 100),
                                "message": message
                            }
                        }
                        print(json.dumps(event, ensure_ascii=False))
                        sys.stdout.flush()

                    cvd_service = CVDService(log_callback, allow_existing_instance=False)
                    result["result"] = cvd_service.extract_pair_cvd(
                        project_id=project_id,
                        pair_id=pair_id,
                        template_path=template_path,
                        filled_path=filled_path,
                        progress_callback=progress_callback
                    )
            except Exception as e:
                print(f"[CVDService] Error: {str(e)}", file=sys.stderr)
                result["result"] = {"success": False, "error": str(e)}

        elif method == "cvd:generateDiff":
            try:
                from services.diff_service import DiffService
                from services.config import config

                project_id = params.get("projectId")
                pair_id = params.get("pairId")
                user_data_path = params.get("userDataPath")

                if not all([project_id, pair_id, user_data_path]):
                    result["result"] = {"success": False, "error": "Missing required parameters"}
                else:
                    if not config.initialized:
                        config.set_user_data_path(user_data_path)

                    diff_service = DiffService()
                    result["result"] = diff_service.generate_diff_for_pair(
                        project_id=project_id,
                        pair_id=pair_id
                    )
            except Exception as e:
                print(f"[DiffService] Error: {str(e)}", file=sys.stderr)
                result["result"] = {"success": False, "error": str(e)}

        elif method == "quit":
            result["result"] = {"success": True}

        else:
            result["error"] = f"Unknown method: {method}"

    except Exception as e:
        result["error"] = str(e)
        result["trace"] = traceback.format_exc()

    return result


def main():
    print("Inserty File Reader Process started", file=sys.stderr)
    sys.stderr.flush()

    try:
        for line in sys.stdin:
            line = line.strip()
            if not line:
                continue

            try:
                request = json.loads(line)
                response = handle_request(request)

                print(json.dumps(response, ensure_ascii=False))
                sys.stdout.flush()

                if request.get("method") == "quit":
                    break

            except json.JSONDecodeError as e:
                error_response = {"error": f"JSON parse error: {e}"}
                print(json.dumps(error_response))
                sys.stdout.flush()

    except KeyboardInterrupt:
        pass


if __name__ == "__main__":
    main()
